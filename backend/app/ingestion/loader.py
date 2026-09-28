import json
import tempfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".json",
}


@dataclass
class RawDocument:
    source_file: str
    text: str


def _load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _load_pdf(path: Path) -> str:
    with pymupdf.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def _load_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _load_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    lines = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
    return "\n".join(lines)


def _load_xlsx(path: Path) -> str:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) for cell in row if cell is not None]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    finally:
        workbook.close()


class _VisibleTextExtractor(HTMLParser):
    """Strips tags, keeping only visible text (skips <script>/<style> bodies)."""

    _SKIP_TAGS = {"script", "style"}

    def __init__(self):
        super().__init__()
        self._skip_depth = 0
        self.chunks: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1

    def handle_endtag(self, tag):
        if tag in self._SKIP_TAGS and self._skip_depth > 0:
            self._skip_depth -= 1

    def handle_data(self, data):
        if self._skip_depth == 0 and data.strip():
            self.chunks.append(data.strip())


def _load_html(path: Path) -> str:
    parser = _VisibleTextExtractor()
    parser.feed(path.read_text(encoding="utf-8", errors="ignore"))
    return "\n".join(parser.chunks)


def _load_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8"))

    def render(value, prefix: str = "") -> list[str]:
        if isinstance(value, dict):
            lines = []
            for key, val in value.items():
                label = f"{prefix}{key}"
                if isinstance(val, (dict, list)):
                    lines.extend(render(val, prefix=f"{label}."))
                else:
                    lines.append(f"{label}: {val}")
            return lines
        if isinstance(value, list):
            lines = []
            for index, item in enumerate(value):
                lines.extend(render(item, prefix=f"{prefix}[{index}]."))
            return lines
        return [f"{prefix.rstrip('.')}: {value}"]

    if isinstance(data, (dict, list)):
        return "\n".join(render(data))
    return str(data)


_LOADERS = {
    ".txt": _load_txt,
    ".md": _load_txt,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".pptx": _load_pptx,
    ".xlsx": _load_xlsx,
    ".html": _load_html,
    ".htm": _load_html,
    ".json": _load_json,
}


def load_document(path: Path) -> RawDocument:
    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise ValueError(f"Unsupported file type: {path.name}")

    return RawDocument(source_file=path.name, text=loader(path))


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from raw upload bytes by dispatching on `filename`'s
    extension, reusing every per-format extractor above via a short-lived
    temp file (several of the underlying libraries need a real path)."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {filename}")

    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(content)
        tmp_path = Path(tmp.name)
    try:
        return load_document(tmp_path).text
    finally:
        tmp_path.unlink(missing_ok=True)


def load_documents(directory: Path, filenames: list[str] | None = None) -> list[RawDocument]:
    """Load documents from a directory.

    If `filenames` is given, only those exact files are loaded (used to enforce
    the "only these 6 new cards" ingestion constraint). Otherwise every
    supported file in the directory is loaded.
    """
    if filenames is not None:
        paths = [directory / name for name in filenames]
        missing = [str(p) for p in paths if not p.exists()]
        if missing:
            raise FileNotFoundError(f"Missing recipe source files: {missing}")
    else:
        paths = sorted(
            p for p in directory.iterdir() if p.suffix.lower() in SUPPORTED_EXTENSIONS
        )

    return [load_document(p) for p in paths]
