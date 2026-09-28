import json

import pytest
from openpyxl import Workbook
from pptx import Presentation

from app.chunking.generic_chunker import chunk_text
from app.ingestion.loader import extract_text
from app.retrieval.retriever import Retriever
from app.services import document_service
from app.services.document_service import DocumentUploadError
from app.vectorstore.chroma_store import ChromaStore


# -- loader.extract_text: one round trip per new format ---------------------


def test_extract_text_txt():
    assert extract_text("notes.txt", b"hello world") == "hello world"


def test_extract_text_html():
    html = b"<html><head><style>.x{}</style></head><body><h1>Title</h1><p>Body text</p></body></html>"
    text = extract_text("page.html", html)
    assert "Title" in text
    assert "Body text" in text
    assert ".x{}" not in text


def test_extract_text_json():
    payload = json.dumps({"name": "Widget", "specs": {"weight": "2kg"}, "tags": ["a", "b"]}).encode()
    text = extract_text("data.json", payload)
    assert "name: Widget" in text
    assert "specs.weight: 2kg" in text
    assert "tags.[0]: a" in text


def test_extract_text_pptx(tmp_path):
    path = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "My Slide Title"
    presentation.save(path)
    text = extract_text("deck.pptx", path.read_bytes())
    assert "My Slide Title" in text


def test_extract_text_xlsx(tmp_path):
    path = tmp_path / "sheet.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Ingredient", "Weight"])
    sheet.append(["Flour", "500g"])
    workbook.save(path)
    text = extract_text("sheet.xlsx", path.read_bytes())
    assert "Flour" in text
    assert "500g" in text


def test_extract_text_unsupported_extension():
    with pytest.raises(ValueError):
        extract_text("archive.zip", b"not real content")


# -- generic_chunker.chunk_text ---------------------------------------------


def test_chunk_text_windows_with_overlap():
    text = "a" * 1000
    pieces = chunk_text(text, chunk_size=300, chunk_overlap=50)
    assert len(pieces) > 1
    assert all(len(p) <= 300 for p in pieces)


def test_chunk_text_empty_returns_no_chunks():
    assert chunk_text("", chunk_size=300, chunk_overlap=50) == []


def test_chunk_text_rejects_invalid_overlap():
    with pytest.raises(ValueError):
        chunk_text("some text", chunk_size=100, chunk_overlap=100)


# -- document_service: upload / list / search / delete round trip ----------


@pytest.fixture
def upload_store(tmp_path_factory):
    return ChromaStore(tmp_path_factory.mktemp("uploaded_chroma_test"))


def test_upload_list_search_delete_round_trip(upload_store, embedding_model):
    content = b"The annual report shows revenue grew by twelve percent this fiscal year."
    info = document_service.upload_document(
        filename="report.txt",
        content=content,
        embedding_model=embedding_model,
        store=upload_store,
        chunk_size=800,
        chunk_overlap=150,
        max_file_size_mb=20,
    )
    assert info.chunk_count == 1
    assert info.source_file == "report.txt"

    listed = document_service.list_documents(upload_store)
    assert any(d.doc_id == info.doc_id for d in listed)

    scoped_retriever = Retriever(upload_store, embedding_model)
    results = scoped_retriever.search_uploaded("revenue growth", top_k=5)
    assert any(r.metadata.doc_id == info.doc_id for r in results)

    deleted = document_service.delete_document(upload_store, info.doc_id)
    assert deleted is True
    assert document_service.list_documents(upload_store) == []
    assert document_service.delete_document(upload_store, info.doc_id) is False


def test_upload_rejects_unsupported_extension(upload_store, embedding_model):
    with pytest.raises(DocumentUploadError):
        document_service.upload_document(
            filename="archive.zip",
            content=b"not real content",
            embedding_model=embedding_model,
            store=upload_store,
            chunk_size=800,
            chunk_overlap=150,
            max_file_size_mb=20,
        )


def test_upload_rejects_oversized_file(upload_store, embedding_model):
    with pytest.raises(DocumentUploadError):
        document_service.upload_document(
            filename="big.txt",
            content=b"x" * 1024,
            embedding_model=embedding_model,
            store=upload_store,
            chunk_size=800,
            chunk_overlap=150,
            max_file_size_mb=0,
        )


def test_upload_rejects_empty_text(upload_store, embedding_model):
    with pytest.raises(DocumentUploadError):
        document_service.upload_document(
            filename="blank.txt",
            content=b"   \n\n  ",
            embedding_model=embedding_model,
            store=upload_store,
            chunk_size=800,
            chunk_overlap=150,
            max_file_size_mb=20,
        )
